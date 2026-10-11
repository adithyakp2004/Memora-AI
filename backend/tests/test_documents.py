from io import BytesIO
from pathlib import Path
from uuid import UUID

import pymupdf
import pytest
from docx import Document as WordDocument
from sqlalchemy import event, select
from sqlalchemy.exc import SQLAlchemyError
from sqlalchemy.orm import Session

from app.db.database import get_engine
from app.main import app
from app.models.document import Document
from app.models.user import User

PASSWORD = "document-password-123"


def register_user(client, email: str) -> tuple[str, UUID]:
    response = client.post(
        "/api/auth/register",
        json={
            "full_name": "Document Owner",
            "email": email,
            "password": PASSWORD,
        },
    )
    assert response.status_code == 201
    user_id = UUID(response.json()["id"])

    login = client.post(
        "/api/auth/login",
        json={"email": email, "password": PASSWORD},
    )
    assert login.status_code == 200
    return login.json()["access_token"], user_id


def pdf_bytes(text: str = "PDF document body.") -> bytes:
    document = pymupdf.open()
    page = document.new_page()
    page.insert_text((72, 72), text)
    content = document.tobytes()
    document.close()
    return content


def docx_bytes() -> bytes:
    document = WordDocument()
    document.add_paragraph("DOCX paragraph content.")
    table = document.add_table(rows=1, cols=2)
    table.cell(0, 0).text = "Table one"
    table.cell(0, 1).text = "Table two"
    stream = BytesIO()
    document.save(stream)
    return stream.getvalue()


def upload(client, token, filename, content, media_type):
    return client.post(
        "/api/documents/upload",
        headers={"Authorization": f"Bearer {token}"},
        files={"file": (filename, content, media_type)},
    )


def test_upload_pdf_extracts_text_and_returns_safe_metadata(client):
    token, _ = register_user(client, "pdf-owner@example.com")
    response = upload(
        client,
        token,
        "notes.pdf",
        pdf_bytes(),
        "application/pdf",
    )

    assert response.status_code == 201
    metadata = response.json()
    assert metadata["file_type"] == ".pdf"
    assert metadata["processing_status"] == "completed"
    assert "extracted_text" not in metadata
    assert "stored_filename" not in metadata
    assert "error_message" in metadata

    content = client.get(
        f"/api/documents/{metadata['id']}/content",
        headers={"Authorization": f"Bearer {token}"},
    )
    assert content.status_code == 200
    assert "PDF document body." in content.json()["extracted_text"]


def test_upload_docx_extracts_paragraphs_and_tables(client):
    token, _ = register_user(client, "docx-owner@example.com")
    response = upload(
        client,
        token,
        "notes.docx",
        docx_bytes(),
        "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
    )

    assert response.status_code == 201
    content = client.get(
        f"/api/documents/{response.json()['id']}/content",
        headers={"Authorization": f"Bearer {token}"},
    )
    assert content.status_code == 200
    assert "DOCX paragraph content." in content.json()["extracted_text"]
    assert "Table one | Table two" in content.json()["extracted_text"]


def test_upload_empty_pdf_page_does_not_claim_ocr_is_required(client):
    token, _ = register_user(client, "scanned-pdf-owner@example.com")
    document = pymupdf.open()
    document.new_page()
    content = document.tobytes()
    document.close()

    response = upload(
        client,
        token,
        "scanned.pdf",
        content,
        "application/pdf",
    )

    assert response.status_code == 422
    assert "OCR" not in response.json()["detail"]
    assert "No selectable text" in response.json()["detail"]
    assert client.get(
        "/api/documents",
        headers={"Authorization": f"Bearer {token}"},
    ).json() == []


def test_upload_txt_supports_utf8_bom_and_normalizes_whitespace(client):
    token, _ = register_user(client, "txt-owner@example.com")
    response = upload(
        client,
        token,
        "notes.txt",
        b"\xef\xbb\xbfFirst   line.\r\n\r\n\r\nSecond line.",
        "text/plain",
    )

    assert response.status_code == 201
    content = client.get(
        f"/api/documents/{response.json()['id']}/content",
        headers={"Authorization": f"Bearer {token}"},
    )
    assert content.json()["extracted_text"] == "First line.\n\nSecond line."


def test_unsupported_extension_and_mismatched_mime_are_rejected(client):
    token, _ = register_user(client, "mime-owner@example.com")

    unsupported = upload(
        client,
        token,
        "script.exe",
        b"not a document",
        "application/octet-stream",
    )
    mismatched = upload(
        client,
        token,
        "notes.pdf",
        pdf_bytes(),
        "text/plain",
    )

    assert unsupported.status_code == 415
    assert mismatched.status_code == 415


def test_server_filename_is_safe_for_a_client_path(client):
    token, _ = register_user(client, "path-owner@example.com")
    response = upload(
        client,
        token,
        r"..\..\private.txt",
        b"Safe storage filename.",
        "text/plain",
    )

    assert response.status_code == 201
    assert response.json()["original_filename"] == "private.txt"
    assert "stored_filename" not in response.json()


def test_empty_and_oversized_files_are_rejected(client, monkeypatch):
    from app.core.config import settings

    token, _ = register_user(client, "size-owner@example.com")
    empty = upload(client, token, "empty.txt", b"", "text/plain")
    monkeypatch.setattr(settings, "document_max_upload_size_mb", 1)
    oversized = upload(
        client,
        token,
        "large.txt",
        b"x" * (1024 * 1024 + 1),
        "text/plain",
    )

    assert empty.status_code == 400
    assert oversized.status_code == 413


def test_extracted_text_limit_is_enforced(client, monkeypatch):
    from app.core.config import settings

    token, _ = register_user(client, "text-limit-owner@example.com")
    monkeypatch.setattr(settings, "document_max_extracted_text_chars", 8)
    response = upload(
        client,
        token,
        "long.txt",
        b"More than eight characters.",
        "text/plain",
    )

    assert response.status_code == 413
    assert client.get(
        "/api/documents",
        headers={"Authorization": f"Bearer {token}"},
    ).json() == []


def test_storage_configuration_cannot_target_frontend(client, monkeypatch):
    from fastapi import HTTPException

    from app.api.routes.documents import _storage_directory
    from app.core.config import BACKEND_DIR, settings

    monkeypatch.setattr(
        settings,
        "document_storage_dir",
        BACKEND_DIR.parent / "frontend" / "public" / "uploads",
    )

    with pytest.raises(HTTPException) as error:
        _storage_directory("test-user")

    assert error.value.status_code == 500


@pytest.mark.parametrize(
    ("filename", "data", "media_type"),
    [
        ("corrupt.pdf", b"%PDF-not-a-real-file", "application/pdf"),
        ("corrupt.docx", b"PK\x03\x04not-a-docx", "application/octet-stream"),
        ("empty-text.txt", b" \r\n\t", "text/plain"),
        ("invalid-encoding.txt", b"\xff\xfe\x81", "text/plain"),
    ],
)
def test_corrupt_or_unextractable_documents_are_rejected(
    client, filename, data, media_type
):
    token, _ = register_user(client, f"invalid-{filename}@example.com")
    response = upload(client, token, filename, data, media_type)
    assert response.status_code in {415, 422}


def test_invalid_upload_leaves_no_file(client):
    from app.core.config import settings

    token, _ = register_user(client, "cleanup-owner@example.com")
    response = upload(
        client,
        token,
        "broken.pdf",
        b"%PDF-invalid",
        "application/pdf",
    )

    assert response.status_code == 422
    files = list(settings.document_storage_dir.rglob("*"))
    assert not any(path.is_file() for path in files)


def test_unauthenticated_upload_and_listing_are_rejected(client):
    response = client.post(
        "/api/documents/upload",
        files={"file": ("notes.txt", b"Private", "text/plain")},
    )
    listing = client.get("/api/documents")

    assert response.status_code == 401
    assert listing.status_code == 401


def test_list_metadata_content_and_document_detail_are_owner_scoped(client):
    owner_token, _ = register_user(client, "owner-scope@example.com")
    other_token, _ = register_user(client, "other-scope@example.com")
    created = upload(
        client,
        owner_token,
        "private.txt",
        b"Private extracted information.",
        "text/plain",
    ).json()
    owner_headers = {"Authorization": f"Bearer {owner_token}"}
    other_headers = {"Authorization": f"Bearer {other_token}"}

    owner_list = client.get("/api/documents", headers=owner_headers)
    other_list = client.get("/api/documents", headers=other_headers)
    owner_detail = client.get(
        f"/api/documents/{created['id']}",
        headers=owner_headers,
    )
    other_detail = client.get(
        f"/api/documents/{created['id']}",
        headers=other_headers,
    )
    other_content = client.get(
        f"/api/documents/{created['id']}/content",
        headers=other_headers,
    )

    assert [item["id"] for item in owner_list.json()] == [created["id"]]
    assert other_list.json() == []
    assert owner_detail.status_code == 200
    assert "extracted_text" not in owner_detail.json()
    assert other_detail.status_code == other_content.status_code == 404
    assert "Private extracted information." not in owner_list.text


def test_delete_is_owner_scoped_and_removes_file(client):
    owner_token, _ = register_user(client, "delete-owner@example.com")
    other_token, _ = register_user(client, "delete-other@example.com")
    created = upload(
        client,
        owner_token,
        "delete.txt",
        b"Delete this document.",
        "text/plain",
    ).json()
    document_id = UUID(created["id"])
    owner_headers = {"Authorization": f"Bearer {owner_token}"}
    other_headers = {"Authorization": f"Bearer {other_token}"}

    with Session(app.state.test_engine) as db:
        document = db.get(Document, document_id)
        stored_path = Path(settings_path()) / str(document.user_id) / document.stored_filename
    assert stored_path.is_file()

    forbidden = client.delete(
        f"/api/documents/{document_id}",
        headers=other_headers,
    )
    assert forbidden.status_code == 404
    assert stored_path.is_file()

    deleted = client.delete(
        f"/api/documents/{document_id}",
        headers=owner_headers,
    )
    assert deleted.status_code == 204
    assert not stored_path.exists()
    assert client.get(
        f"/api/documents/{document_id}",
        headers=owner_headers,
    ).status_code == 404


def settings_path():
    from app.core.config import settings

    return settings.document_storage_dir


def test_upload_database_failure_removes_stored_file(client):
    from app.core.config import settings

    token, _ = register_user(client, "db-failure-owner@example.com")

    def fail_commit(_session):
        raise SQLAlchemyError("test failure")

    event.listen(Session, "before_commit", fail_commit, once=True)
    try:
        response = upload(
            client,
            token,
            "rollback.txt",
            b"This must not leave an orphan.",
            "text/plain",
        )
    finally:
        event.remove(Session, "before_commit", fail_commit)

    assert response.status_code == 500
    assert not any(
        path.is_file() for path in settings.document_storage_dir.rglob("*")
    )
