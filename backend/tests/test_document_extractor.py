import zipfile
from io import BytesIO
from pathlib import Path

import pymupdf
import pytest
from docx import Document as WordDocument

from app.services.document_extractor import (
    DocumentProcessingError,
    extract_text,
)


def create_pdf(*page_texts: str, encryption: bool = False) -> bytes:
    document = pymupdf.open()
    for text in page_texts:
        page = document.new_page()
        if text:
            page.insert_text((72, 72), text)
    save_options = {}
    if encryption:
        save_options = {
            "encryption": pymupdf.PDF_ENCRYPT_AES_256,
            "owner_pw": "synthetic-owner-password",
            "user_pw": "synthetic-user-password",
        }
    content = document.tobytes(**save_options)
    document.close()
    return content


def create_image_pdf() -> bytes:
    pixmap = pymupdf.Pixmap(
        pymupdf.csRGB,
        pymupdf.IRect(0, 0, 64, 64),
        0,
    )
    pixmap.clear_with(220)
    document = pymupdf.open()
    page = document.new_page()
    page.insert_image(page.rect, pixmap=pixmap)
    content = document.tobytes()
    document.close()
    return content


def create_docx(*paragraphs: str, table: bool = False) -> bytes:
    document = WordDocument()
    for text in paragraphs:
        document.add_paragraph(text)
    if table:
        word_table = document.add_table(rows=1, cols=2)
        word_table.cell(0, 0).text = "First cell"
        word_table.cell(0, 1).text = "Second cell"
    stream = BytesIO()
    document.save(stream)
    return stream.getvalue()


def write_upload(tmp_path: Path, filename: str, content: bytes) -> Path:
    path = tmp_path / filename
    path.write_bytes(content)
    return path


@pytest.mark.parametrize(
    ("page_texts", "expected"),
    [
        (("Known PDF text.",), ("Known PDF text.",)),
        (
            ("First PDF page.", "Second PDF page."),
            ("First PDF page.", "Second PDF page."),
        ),
        (
            ("", "Text after an empty page."),
            ("Text after an empty page.",),
        ),
    ],
)
def test_pdf_extracts_text_in_page_order_and_skips_empty_pages(
    tmp_path, page_texts, expected
):
    path = write_upload(tmp_path, "text.pdf", create_pdf(*page_texts))

    extracted = extract_text(path, ".pdf", 1_000_000)

    for text in expected:
        assert text in extracted
    positions = [extracted.index(text) for text in expected]
    assert positions == sorted(positions)


def test_empty_pdf_page_does_not_get_an_ocr_required_classification(tmp_path, caplog):
    path = write_upload(tmp_path, "blank.pdf", create_pdf(""))
    caplog.set_level("INFO", logger="app.services.document_extractor")

    with pytest.raises(DocumentProcessingError) as error:
        extract_text(path, ".pdf", 1_000_000)

    assert error.value.category == "pdf_no_selectable_text"
    assert "OCR" not in str(error.value)
    assert "page_count=1" in caplog.text
    assert "raw_char_count=" in caplog.text
    assert "image_count=0" in caplog.text
    assert "error_category=pdf_no_selectable_text" in caplog.text
    assert "outcome=failed" in caplog.text
    assert "blank.pdf" not in caplog.text


def test_full_page_image_pdf_is_only_labeled_likely_image_based(tmp_path):
    path = write_upload(tmp_path, "image-only.pdf", create_image_pdf())

    with pytest.raises(DocumentProcessingError) as error:
        extract_text(path, ".pdf", 1_000_000)

    assert error.value.category == "pdf_likely_image_based"
    assert "appears to contain page images" in str(error.value)
    assert "OCR is not available" in str(error.value)


def test_pdf_text_drawings_that_decode_as_whitespace_are_not_called_image_only(
    tmp_path, monkeypatch
):
    path = write_upload(tmp_path, "unmapped-text.pdf", create_pdf("Visible text."))
    monkeypatch.setattr(pymupdf.Page, "get_text", lambda self, *args, **kwargs: " \n")

    with pytest.raises(DocumentProcessingError) as error:
        extract_text(path, ".pdf", 1_000_000)

    assert error.value.category == "pdf_text_decode_failed"
    assert "text elements" in str(error.value)
    assert "OCR" not in str(error.value)


@pytest.mark.parametrize(
    "content",
    [
        b"%PDF-this is not a valid PDF",
        b"",
    ],
)
def test_corrupt_or_empty_pdf_is_reported_safely(tmp_path, content):
    path = write_upload(tmp_path, "corrupt.pdf", content)

    with pytest.raises(DocumentProcessingError):
        extract_text(path, ".pdf", 1_000_000)


def test_encrypted_pdf_is_rejected_without_password_prompt(tmp_path):
    path = write_upload(
        tmp_path,
        "encrypted.pdf",
        create_pdf("Protected content.", encryption=True),
    )

    with pytest.raises(DocumentProcessingError) as error:
        extract_text(path, ".pdf", 1_000_000)

    assert error.value.category == "encrypted_pdf"
    assert "Password-protected" in str(error.value)


def test_docx_extracts_multiple_paragraphs_and_table_text(tmp_path):
    path = write_upload(
        tmp_path,
        "notes.docx",
        create_docx("First paragraph.", "", "Third paragraph.", table=True),
    )

    extracted = extract_text(path, ".docx", 1_000_000)

    assert extracted.index("First paragraph.") < extracted.index("Third paragraph.")
    assert "First cell | Second cell" in extracted


def test_empty_docx_has_no_extractable_text(tmp_path):
    path = write_upload(tmp_path, "empty.docx", create_docx())

    with pytest.raises(DocumentProcessingError) as error:
        extract_text(path, ".docx", 1_000_000)

    assert error.value.category == "no_text"
    assert "No extractable text" in str(error.value)


@pytest.mark.parametrize(
    "content",
    [
        b"not a ZIP archive",
        b"PK\x03\x04invalid ZIP data",
    ],
)
def test_corrupt_docx_is_reported_safely(tmp_path, content):
    path = write_upload(tmp_path, "corrupt.docx", content)

    with pytest.raises(DocumentProcessingError):
        extract_text(path, ".docx", 1_000_000)


def test_invalid_zip_without_docx_document_is_rejected(tmp_path):
    stream = BytesIO()
    with zipfile.ZipFile(stream, "w") as archive:
        archive.writestr("not-a-docx.xml", "<document/>")
    path = write_upload(tmp_path, "renamed.docx", stream.getvalue())

    with pytest.raises(DocumentProcessingError):
        extract_text(path, ".docx", 1_000_000)


@pytest.mark.parametrize(
    ("content", "expected"),
    [
        ("UTF-8 text ✓".encode("utf-8"), "UTF-8 text ✓"),
        (b"\xef\xbb\xbfText with BOM.", "Text with BOM."),
    ],
)
def test_txt_extracts_utf8_and_utf8_bom(tmp_path, content, expected):
    path = write_upload(tmp_path, "notes.txt", content)

    assert extract_text(path, ".txt", 1_000_000) == expected


@pytest.mark.parametrize("content", [b"", b"\xff\xfe\x81"])
def test_empty_and_invalid_txt_are_rejected(tmp_path, content):
    path = write_upload(tmp_path, "invalid.txt", content)

    with pytest.raises(DocumentProcessingError):
        extract_text(path, ".txt", 1_000_000)


def test_diagnostics_include_metadata_but_never_extracted_text(
    tmp_path, caplog
):
    private_text = "Synthetic private document content."
    path = write_upload(tmp_path, "private.txt", private_text.encode("utf-8"))
    caplog.set_level("INFO", logger="app.services.document_extractor")

    assert extract_text(path, ".txt", 1_000_000) == private_text

    logs = caplog.text
    assert "file_type=.txt" in logs
    assert f"file_size_bytes={len(private_text.encode('utf-8'))}" in logs
    assert "method=utf-8" in logs
    assert "extracted_char_count=" in logs
    assert "outcome=completed" in logs
    assert private_text not in logs
